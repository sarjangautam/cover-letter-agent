from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt


def create_cover_letter_docx(
    output_path: Path,
    candidate,
    job,
    cover_letter,
) -> None:
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)

    header_lines = [
        candidate.name,
        candidate.address,
        candidate.phone,
        candidate.email,
        date.today().strftime("%d %B %Y"),
        "Hiring Manager",
        job.company,
    ]

    for index, line in enumerate(header_lines):
        if line and line.strip():
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(line)

        if index in (3, 4):
            document.add_paragraph().paragraph_format.space_after = Pt(0)

    document.add_paragraph().paragraph_format.space_after = Pt(4)

    greeting = document.add_paragraph(cover_letter.greeting)
    greeting.paragraph_format.space_after = Pt(10)

    opening = document.add_paragraph(cover_letter.opening)
    opening.paragraph_format.space_after = Pt(8)

    for paragraph_text in cover_letter.body:
        paragraph = document.add_paragraph(paragraph_text)
        paragraph.paragraph_format.space_after = Pt(8)

    closing = document.add_paragraph(cover_letter.closing)
    closing.paragraph_format.space_before = Pt(6)
    closing.paragraph_format.space_after = Pt(4)

    sign_off = cover_letter.sign_off.strip()
    candidate_name = candidate.name.strip()

    if sign_off.casefold().endswith(candidate_name.casefold()):
        sign_off = sign_off[: -len(candidate_name)].rstrip()

    signoff_paragraph = document.add_paragraph(sign_off)
    signoff_paragraph.paragraph_format.space_after = Pt(2)

    name_paragraph = document.add_paragraph(candidate_name)
    name_paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    document.save(output_path)
