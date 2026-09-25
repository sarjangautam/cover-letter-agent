from pathlib import Path

from docx import Document


def read_docx(file_path: Path) -> str:
    document = Document(file_path)

    content = []

    # Read paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            content.append(text)

    # Read tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    row_text.append(text)

            if row_text:
                content.append(" | ".join(row_text))

    return "\n".join(content)


if __name__ == "__main__":

    resume_folder = Path("Resume")

    resume_files = list(resume_folder.glob("*.docx"))

    if not resume_files:
        print("No DOCX resume found in the Resume folder.")
        exit()

    if len(resume_files) > 1:
        print("Multiple resume files found:")

        for file in resume_files:
            print(f"  - {file.name}")

        print("\nPlease keep only one resume for now.")
        exit()

    resume_file = resume_files[0]

    print(f"Reading resume: {resume_file.name}")
    print("=" * 80)

    resume_text = read_docx(resume_file)

    print(resume_text)